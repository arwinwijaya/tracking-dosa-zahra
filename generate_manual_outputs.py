"""Generate all transaction outputs exclusively from manual.txt; does not access images/OCR."""
from pathlib import Path
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from src.manual_loader import load_manual_txt, validate_running_balance
from src import deduplication, reconciliation, classification, fraud_analysis
from src.statement_totals import TOTALS
from src.utils import money, mask
from src import config


def main():
    out = config.OUTPUT
    out.mkdir(parents=True, exist_ok=True)
    df = load_manual_txt(config.ROOT / 'manual.txt')
    if len(df) != 86:
        raise ValueError(f'Expected 86 manual rows; parsed {len(df)}')
    if df[['debet', 'kredit', 'saldo']].isna().any().any():
        raise ValueError('Manual table contains an unparsed debit/credit/balance value')
    errors = validate_running_balance(df)
    # Validate each month against the manually transcribed statement opening/closing controls.
    dt = pd.to_datetime(df.tanggal_transaksi, format='%d/%m/%y')
    month_reports = {}
    for month, control in TOTALS.items():
        group = df.loc[dt.dt.strftime('%Y-%m') == month]
        if group.empty:
            raise ValueError(f'Missing manual rows for {month}')
        debit, credit = int(group.debet.sum()), int(group.kredit.sum())
        first_expected = control['opening_balance'] - int(group.iloc[0].debet) + int(group.iloc[0].kredit)
        if first_expected != int(group.iloc[0].saldo):
            raise ValueError(f'{month} opening-to-first-transaction balance mismatch: expected {first_expected}, got {group.iloc[0].saldo}')
        if debit != control['total_debit'] or credit != control['total_credit']:
            raise ValueError(f'{month} manual totals differ from controls: debit {debit}/{control["total_debit"]}, credit {credit}/{control["total_credit"]}')
        if control['opening_balance'] - debit + credit != control['closing_balance']:
            raise ValueError(f'{month} monthly control balance does not reconcile')
        month_reports[month] = {
            'opening_balance': control['opening_balance'], 'total_debit_manual': debit,
            'total_credit_manual': credit, 'closing_balance_expected': control['closing_balance'],
            'closing_balance_manual': int(group.iloc[-1].saldo),
            'closing_balance_match': int(group.iloc[-1].saldo) == control['closing_balance'],
            'transactions': len(group),
        }
    if errors:
        raise ValueError(f'{len(errors)} running balance mismatches: {errors[:3]}')

    df = deduplication.mark(df)
    df = reconciliation.reconcile(df)
    if (df.reconciliation_status != 'MATCH').any():
        raise ValueError('Sequential saldo reconciliation produced non-MATCH rows')
    df = classification.classify(df)
    df = fraud_analysis.assess(df)
    df.to_csv(out / 'transactions_audit.csv', index=False, encoding='utf-8-sig')
    df.to_csv(out / 'manual_transactions.csv', index=False, encoding='utf-8-sig')
    (out / 'raw_transactions.json').write_text(df.to_json(orient='records', indent=2, force_ascii=False), encoding='utf-8')

    monthly = df.assign(month=pd.to_datetime(df.tanggal_transaksi, format='%d/%m/%y').dt.strftime('%Y-%m')).groupby('month').agg(
        total_debit=('debet','sum'), total_credit=('kredit','sum'), transaction_count=('debet','size'), closing_balance=('saldo','last')).reset_index()
    monthly['opening_balance'] = monthly.month.map(lambda m: TOTALS[m]['opening_balance'])
    monthly['statement_closing_balance'] = monthly.month.map(lambda m: TOTALS[m]['closing_balance'])
    monthly['net_cash_flow'] = monthly.total_credit - monthly.total_debit
    recon = df[['tanggal_transaksi','waktu','debet','kredit','expected_saldo','saldo','saldo_difference','reconciliation_status','uraian_transaksi','teller']].copy()
    recon.insert(0,'row',range(1,len(recon)+1))
    summary = pd.DataFrame([
        ['Source','manual.txt only'], ['Transaction count',len(df)],
        ['Period',f'{dt.min().date()} — {dt.max().date()}'], ['Total debit',int(df.debet.sum())],
        ['Total credit',int(df.kredit.sum())], ['Opening balance',TOTALS['2026-07']['opening_balance']],
        ['Closing balance',int(df.saldo.iloc[-1])], ['Running balance mismatches',len(errors)],
        ['Sequential reconciliation', 'PASS' if (df.reconciliation_status=='MATCH').all() else 'FAIL'],
        ['Statement monthly totals', 'PASS' if all(v['closing_balance_match'] for v in month_reports.values()) else 'FAIL'],
    ], columns=['Metric','Value'])
    books = {
        'Ringkasan': summary,
        'Transaksi Manual': df,
        'Rekonsiliasi Saldo': recon,
        'Rekap Bulanan': monthly,
        'Kontrol Saldo Bulanan': pd.DataFrame.from_dict(month_reports, orient='index').rename_axis('month').reset_index(),
    }
    for path in (config.EXCEL, out/'manual_BRI_Analysis.xlsx'):
        try:
            if path.exists(): path.unlink()
        except PermissionError:
            print(f"Warning: Could not overwrite {path} because it is open in Excel. Skipping.")
            continue
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            for sheet, data in books.items(): data.to_excel(writer, sheet_name=sheet, index=False)
        wb=load_workbook(path)
        for ws in wb.worksheets:
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='17365D')
            for col in ws.columns:
                letter=get_column_letter(col[0].column)
                ws.column_dimensions[letter].width=min(60,max(12,max((len(str(c.value)) for c in col if c.value is not None),default=0)+2))
            headers={c.value:c.column for c in ws[1]}
            for name,j in headers.items():
                if str(name).lower() in {'debet','kredit','saldo','expected_saldo','saldo_difference','total_debit','total_credit','opening_balance','closing_balance','statement_closing_balance','closing_balance_expected','closing_balance_manual','net_cash_flow'}:
                    for cells in ws.iter_rows(min_row=2,min_col=j,max_col=j): cells[0].number_format='"Rp "#,##0;[Red]("Rp "#,##0)'
        wb.save(path)

    lines=['# Analisis Rekening BRI — Berdasarkan Data Manual','','**Sumber transaksi: `manual.txt` saja. Tidak ada gambar atau hasil OCR yang digunakan untuk menyusun keluaran ini.**','',
           '## Ringkasan',f'- Periode: {dt.min().date()} — {dt.max().date()}',f'- Jumlah transaksi: {len(df)}',
           f'- Total debet: {money(df.debet.sum())}',f'- Total kredit: {money(df.kredit.sum())}',
           f'- Saldo akhir: {money(df.saldo.iloc[-1])}',f'- Rekonsiliasi berurutan debet/kredit/saldo: **LULUS ({len(df)} baris cocok)**.','',
           '## Validasi bulanan']
    for month, v in month_reports.items():
        lines.append(f"- {month}: debit {money(v['total_debit_manual'])}, kredit {money(v['total_credit_manual'])}, saldo akhir manual {money(v['closing_balance_manual'])} vs kontrol {money(v['closing_balance_expected'])} — **{'COCOK' if v['closing_balance_match'] else 'TIDAK COCOK'}**.")
    lines += ['', 'Nilai transaksi, termasuk deskripsi/teller dan saldo, mengikuti transkripsi manual. Tidak ada nominal yang diubah untuk menyesuaikan total kontrol. Keluaran ini adalah analisis transaksi, bukan penetapan tindak pidana.']
    for path in (config.REPORT, out/'manual_report.md'):
        path.write_text('\n'.join(lines)+'\n', encoding='utf-8')
    qc={
        'source':'manual.txt only','images_read':False,'images_processed':0,'transactions':len(df),
        'total_debit':int(df.debet.sum()),'total_credit':int(df.kredit.sum()),
        'running_balance_mismatches':len(errors),'reconciliation_matches':int((df.reconciliation_status=='MATCH').sum()),
        'reconciliation_mismatches':int((df.reconciliation_status=='MISMATCH').sum()),
        'cannot_validate':int((df.reconciliation_status=='CANNOT_VALIDATE').sum()),
        'monthly_controls':month_reports,
    }
    (out/'quality_control.json').write_text(json.dumps(qc,indent=2,ensure_ascii=False),encoding='utf-8')
    (out/'manual_quality_control.json').write_text(json.dumps(qc,indent=2,ensure_ascii=False),encoding='utf-8')
    print(json.dumps(qc,indent=2,ensure_ascii=False))
    print('Updated:',config.EXCEL,config.REPORT,out/'transactions_audit.csv')

if __name__=='__main__': main()
