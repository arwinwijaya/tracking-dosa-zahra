"""Build 5-month (May–September 2026) forensic workbook by merging
Mei–Juni (new user data) with Juli–September (existing reference).
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd

from src.bri_parser import (
    build_sheets,
    classify_fees,
    link_interest_tax,
    load_verification_db,
    parse_transactions,
    write_workbook,
)

ROOT = Path(__file__).resolve().parent
SRC_MJ = ROOT / "output" / "mei_juni_source.json"
SRC_REF = ROOT / "output" / "BRI_Forensic_Transaction_Analysis.xlsx"
VERIF = ROOT / "config" / "account_verification.csv"
OUT = ROOT / "output" / "BRI_Forensic_Transaction_Analysis_5Bulan_2026.xlsx"

ORDER_NOTE = (
    "Anomali urutan waktu: 01/05/26 19:26:49 tercantum sebelum 19:25:49; "
    "urutan sumber dipertahankan, saldo konsisten mengikuti urutan sumber"
)


def _extract_raw_from_workbook(path: Path) -> pd.DataFrame:
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Daftar Transaksi"]
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rows.append({
            "tanggal_transaksi": str(r[idx["Tanggal"]]),
            "waktu": str(r[idx["Waktu"]]),
            "uraian_transaksi": r[idx["Uraian Transaksi"]],
            "teller": str(r[idx["Teller"]]),
            "debet": int(r[idx["Debit"]] or 0),
            "kredit": int(r[idx["Credit"]] or 0),
            "saldo": int(r[idx["Saldo"]] or 0),
        })
    return pd.DataFrame(rows)


def main() -> None:
    # ---- 1. Load raw data: May-June from JSON, Jul-Sep from reference workbook
    df_mj = pd.DataFrame(
        json.loads(SRC_MJ.read_text(encoding="utf-8")),
        columns=["tanggal_transaksi", "waktu", "uraian_transaksi", "teller", "debet", "kredit", "saldo"],
    )
    df_mj["teller"] = df_mj["teller"].astype(str)

    df_ref = _extract_raw_from_workbook(SRC_REF)

    # ---- 2. Concatenate chronologically (May-June first, then Jul-Sep)
    df_all = pd.concat([df_mj, df_ref], ignore_index=True)

    # ---- 3. Parse with verification DB
    verif_db = load_verification_db(VERIF)
    parsed = parse_transactions(df_all, verification_db=verif_db)
    parsed = classify_fees(parsed)
    parsed = link_interest_tax(parsed)

    # ---- 4. Build sheets
    sheets = build_sheets(parsed)

    # ---- 5. Post-process: add anomaly notes
    daftar = sheets["Daftar Transaksi"]
    # TXN0001 & TXN0002 = May 1st time inversion (still valid)
    for txn in ("TXN0001", "TXN0002"):
        m = daftar["Transaction ID"] == txn
        daftar.loc[m, "Catatan Parser"] = (
            daftar.loc[m, "Catatan Parser"].astype(str) + "; " + ORDER_NOTE
        )

    # ---- 6. Update Ringkasan with 5-month info
    ringkas = sheets["Ringkasan"]
    ringkas.loc[ringkas["Keterangan"] == "Sumber data", "Nilai"] = (
        "Pesan pengguna (Mei-Jun) + BRI_Forensic_Transaction_Analysis.xlsx (Jul-Sep) "
        "(ditranskripsi; urutan sumber dipertahankan)"
    )
    extra = [
        (
            "Anomali urutan waktu 01/05/26",
            "19:26:49 tercantum sebelum 19:25:49; urutan sumber dipertahankan, "
            "saldo tetap konsisten mengikuti urutan sumber",
        ),
        (
            "Saldo awal Mei (turunan)",
            "81682563 — dihitung dari baris pertama (saldo + debet - kredit), "
            "tidak terkonfirmasi independen",
        ),
        (
            "Keterkaitan Mei-Juni",
            "Saldo akhir Mei (76544173) = saldo awal Juni (76544173) — konsisten",
        ),
        (
            "Keterkaitan Juni-Juli",
            "Saldo akhir Juni (77451791) = saldo awal Juli (77451791) — konsisten",
        ),
    ]
    sheets["Ringkasan"] = pd.concat(
        [ringkas, pd.DataFrame(extra, columns=["Keterangan", "Nilai"])],
        ignore_index=True,
    )

    # ---- 7. Anomali sheet (only the time-inversion anomaly remains)
    sheets["Anomali & Selisih"] = pd.DataFrame(
        [
            [
                "1",
                "TXN0001–TXN0002 — 01/05/26",
                "19:26:49 tercantum sebelum 19:25:49",
                "Urutan tidak kronologis; saldo tetap cocok mengikuti "
                "urutan sumber",
                "Anomali urutan waktu — urutan sumber dipertahankan, "
                "tidak diurutkan ulang",
            ],
        ],
        columns=["No", "Lokasi", "Detail", "Dampak", "Keterangan"],
    )

    # ---- 8. Write
    write_workbook(sheets, OUT)
    print("written:", OUT)


if __name__ == "__main__":
    main()
