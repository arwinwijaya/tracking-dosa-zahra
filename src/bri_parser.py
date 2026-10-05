"""Evidence-aware parser for BRI statement `Uraian Transaksi` descriptions.

Design rules (see project knowledge base):
  * Raw statement text is the source of truth and is never modified.
  * Nothing is fabricated: unknown names/accounts/merchants stay UNKNOWN.
  * Every interpretation carries a confidence level and a provenance type.
  * Technical transaction codes and economic interpretations are kept separate.
  * Account numbers extracted from WBNK payloads are *candidates* until they
    appear in config/account_verification.csv.

The parser is deliberately reusable: feed it any DataFrame that has a
description column (default `uraian_transaksi`) plus optional debet/kredit.
"""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd

# --------------------------------------------------------------------------
# Provenance vocabulary
# --------------------------------------------------------------------------
OFFICIAL_BRI = "OFFICIAL_BRI"
EXPLICIT_DESCRIPTION = "EXPLICIT_DESCRIPTION"
SUPPORTED_REFERENCE = "SUPPORTED_REFERENCE"
STRUCTURAL_INFERENCE = "STRUCTURAL_INFERENCE"
MANUAL_VERIFICATION = "MANUAL_VERIFICATION"
UNKNOWN_SOURCE = "UNKNOWN"

# Evidence levels
VERIFIED = "VERIFIED"
SUPPORTED = "SUPPORTED"
INFERRED = "INFERRED"
UNKNOWN_EVIDENCE = "UNKNOWN"

UNKNOWN = "UNKNOWN"

# --------------------------------------------------------------------------
# Pattern registry.  Priority order matters: the first match wins.
# --------------------------------------------------------------------------
# NOTE on WBNKTRIK: the reference material associates WBNKTRIK with BRILink
# agent cash withdrawal, while the structural `TRF` sibling is a transfer.
# We therefore record the statement-structural type AND a separate cautious
# economic interpretation instead of collapsing the two.
PATTERN_REGISTRY: dict[str, dict] = {
    "WBNKTRIK": {
        "family": "WBNK",
        "channel": "Web Banking / Electronic Banking",
        "type": "Transfer",
        "economic": "BRILink agent transaction - cash withdrawal candidate (references) vs transfer (statement structure); unresolved",
        "code_meaning": "Public references associate WBNKTRIK with cash withdrawal via Agen BRILink; do not assert a plain account transfer",
        "code_meaning_confidence": "MEDIUM",
        "source_type": SUPPORTED_REFERENCE,
        "evidence": SUPPORTED,
        "account_role": "AGENT_OR_COUNTERPARTY",
    },
    "WBNKTRF": {
        "family": "WBNK",
        "channel": "Web Banking / Electronic Banking",
        "type": "Transfer",
        "economic": "Funds transfer via WBNK/BRILink channel",
        "code_meaning": "Transfer transaction using WBNK/BRILink channel (TRF = transfer indicator)",
        "code_meaning_confidence": "MEDIUM",
        "source_type": STRUCTURAL_INFERENCE,
        "evidence": INFERRED,
        "account_role": "BENEFICIARY",
    },
    "DANA": {
        "family": "E_WALLET",
        "channel": "DANA",
        "type": "Digital Wallet",
        "economic": "DANA wallet transaction (direction from Debit/Credit)",
        "code_meaning": "Transaction explicitly associated with DANA digital wallet/payment infrastructure",
        "code_meaning_confidence": "HIGH",
        "source_type": EXPLICIT_DESCRIPTION,
        "evidence": SUPPORTED,
        "account_role": UNKNOWN,
    },
    "PRCH": {
        "family": "CARD_PURCHASE",
        "channel": "Card / EDC",
        "type": "Purchase",
        "economic": "Merchant purchase (card/EDC)",
        "code_meaning": "Purchase transaction at a merchant, commonly via BRI debit card/EDC",
        "code_meaning_confidence": "HIGH",
        "source_type": SUPPORTED_REFERENCE,
        "evidence": SUPPORTED,
        "account_role": UNKNOWN,
    },
    "PRTT": {
        "family": "RETAIL_TRANSACTION",
        "channel": "BRI_EDC",
        "type": "Retail / Purchase",
        "economic": "Retail purchase at BRI EDC merchant",
        "code_meaning": "Belanja di EDC merchant BRI",
        "code_meaning_confidence": "MEDIUM_HIGH",
        "source_type": SUPPORTED_REFERENCE,
        "evidence": SUPPORTED,
        "account_role": UNKNOWN,
    },
    "NP": {
        "family": "PAYMENT",
        "channel": "EDC / National Payment",
        "type": "Payment",
        "economic": "Merchant payment through National Payment / EDC infrastructure",
        "code_meaning": "National Payment transaction, generally merchant payment through EDC",
        "code_meaning_confidence": "HIGH",
        "source_type": SUPPORTED_REFERENCE,
        "evidence": SUPPORTED,
        "account_role": UNKNOWN,
    },
    "INTEREST": {
        "family": "BANK_SYSTEM",
        "channel": "Bank System",
        "type": "Interest",
        "economic": "Interest credited by the bank",
        "code_meaning": "Interest credited by the bank to the account",
        "code_meaning_confidence": "HIGH",
        "source_type": EXPLICIT_DESCRIPTION,
        "evidence": VERIFIED,
        "account_role": UNKNOWN,
    },
    "TAX": {
        "family": "BANK_SYSTEM",
        "channel": "Bank System",
        "type": "Tax",
        "economic": "Tax deduction",
        "code_meaning": "Tax deduction",
        "code_meaning_confidence": "HIGH",
        "source_type": EXPLICIT_DESCRIPTION,
        "evidence": VERIFIED,
        "account_role": UNKNOWN,
    },
    "ADMIN_FEE": {
        "family": "BANK_FEE",
        "channel": "Bank System",
        "type": "Administration Fee",
        "economic": "Periodic account administration fee",
        "code_meaning": "Periodic account administration fee charged by BRI",
        "code_meaning_confidence": "HIGH",
        "source_type": OFFICIAL_BRI,
        "evidence": VERIFIED,
        "account_role": UNKNOWN,
    },
    "MONTHLY_FEE_ATM": {
        "family": "BANK_FEE",
        "channel": "Bank System",
        "type": "ATM Monthly Fee",
        "economic": "Debit card / ATM monthly administration fee",
        "code_meaning": "Monthly administration/maintenance fee for the BRI debit/ATM card",
        "code_meaning_confidence": "HIGH",
        "source_type": OFFICIAL_BRI,
        "evidence": VERIFIED,
        "account_role": UNKNOWN,
    },
    # --- explicit patterns detected beyond the requested taxonomy ---------
    "BFST": {
        "family": "TRANSFER",
        "channel": "BI-FAST",
        "type": "Transfer",
        "economic": "BI-FAST transfer (direction from Debit/Credit)",
        "code_meaning": "Explicit BI-FAST transfer marker (BFST / FASTIDJA) in the description",
        "code_meaning_confidence": "MEDIUM",
        "source_type": EXPLICIT_DESCRIPTION,
        "evidence": SUPPORTED,
        "account_role": "SENDER",
    },
    "TARIK_TUNAI": {
        "family": "CASH_WITHDRAWAL",
        "channel": "ATM",
        "type": "Cash Withdrawal",
        "economic": "ATM cash withdrawal",
        "code_meaning": "Explicit 'TARIK TUNAI' (cash withdrawal) with ATM ESB reference",
        "code_meaning_confidence": "HIGH",
        "source_type": EXPLICIT_DESCRIPTION,
        "evidence": SUPPORTED,
        "account_role": UNKNOWN,
    },
    "UNKNOWN": {
        "family": UNKNOWN,
        "channel": UNKNOWN,
        "type": UNKNOWN,
        "economic": UNKNOWN,
        "code_meaning": "Pattern not recognised; requires manual review",
        "code_meaning_confidence": "UNKNOWN",
        "source_type": UNKNOWN_SOURCE,
        "evidence": UNKNOWN_EVIDENCE,
        "account_role": UNKNOWN,
    },
}

# Order in which patterns are tested.  Unknown is always last.
PATTERN_PRIORITY = [
    "WBNKTRIK", "WBNKTRF", "DANA", "PRCH", "PRTT", "NP",
    "INTEREST", "TAX", "ADMIN_FEE", "MONTHLY_FEE_ATM",
    "BFST", "TARIK_TUNAI", "UNKNOWN",
]

WBNK_PATTERNS = {"WBNKTRIK", "WBNKTRF"}

ESB_RE = re.compile(r"ESB:([A-Z0-9_]+):([A-Z0-9]+):([A-Z0-9]+)")
WBNK_RE = re.compile(r"^(WBNKTRIK|WBNKTRF)(\d+)T(\d{16})")


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def normalize_description(raw: str) -> str:
    """Uppercase, collapse whitespace and line breaks.  Digits are preserved."""
    return re.sub(r"\s+", " ", str(raw).replace("\n", " ").replace("\r", " ")).strip().upper()


def load_verification_db(path: Path) -> dict[str, dict]:
    """Load manually verified accounts.  Missing file => empty database."""
    if not path.exists():
        return {}
    db = pd.read_csv(path, dtype=str).fillna("")
    out = {}
    for _, r in db.iterrows():
        acct = str(r.get("account_number", "")).strip()
        if acct:
            out[acct] = {
                "account_name": r.get("account_name", UNKNOWN) or UNKNOWN,
                "bank_name": r.get("bank_name", UNKNOWN) or UNKNOWN,
                "verification_method": r.get("verification_method", "") or "",
            }
    return out


def _direction(debet, kredit) -> str:
    d = 0 if pd.isna(debet) else int(debet)
    c = 0 if pd.isna(kredit) else int(kredit)
    if d > 0 and c == 0:
        return "OUT"
    if c > 0 and d == 0:
        return "IN"
    if d == 0 and c == 0:
        return "UNKNOWN"
    return "REVIEW"


def _blank_row() -> dict:
    return dict(
        transaction_pattern=UNKNOWN,
        technical_transaction_code=UNKNOWN,
        transaction_family=UNKNOWN,
        transaction_channel=UNKNOWN,
        transaction_type=UNKNOWN,
        economic_transaction_type=UNKNOWN,
        transaction_direction=UNKNOWN,
        code_meaning=UNKNOWN,
        code_meaning_confidence=UNKNOWN,
        interpretation_source_type=UNKNOWN_SOURCE,
        evidence_level=UNKNOWN_EVIDENCE,
        parser_confidence=UNKNOWN,
        internal_id=UNKNOWN,
        raw_destination_segment=UNKNOWN,
        destination_account_candidate=UNKNOWN,
        technical_counterparty_account_candidate=UNKNOWN,
        account_role=UNKNOWN,
        destination_account_verified=False,
        destination_name=UNKNOWN,
        destination_bank=UNKNOWN,
        counterparty_name=UNKNOWN,
        merchant_name=UNKNOWN,
        wallet_name=UNKNOWN,
        transaction_reference=UNKNOWN,
        secondary_reference=UNKNOWN,
        long_reference_payload=UNKNOWN,
        esb_system=UNKNOWN,
        esb_code=UNKNOWN,
        esb_reference=UNKNOWN,
        is_fee_transaction=False,
        fee_classification=UNKNOWN,
        fee_amount=0,
        parent_transaction_id="",
        possible_paired_transaction=False,
        tax_relationship_candidate=UNKNOWN,
        related_interest_amount="",
        effective_tax_rate="",
        tax_relationship_confidence="",
        manual_review_required=False,
        parser_notes="",
    )


def detect_pattern(norm: str) -> str:
    """Return the first matching pattern in priority order."""
    for pat in PATTERN_PRIORITY:
        if pat == "UNKNOWN":
            continue
        if pat == "WBNKTRIK" and norm.startswith("WBNKTRIK"):
            return pat
        if pat == "WBNKTRF" and norm.startswith("WBNKTRF"):
            return pat
        if pat == "DANA" and norm.startswith("DANA"):
            return pat
        if pat == "PRCH" and norm.startswith("PRCH"):
            return pat
        if pat == "PRTT" and norm.startswith("PRTT"):
            return pat
        if pat == "NP" and re.match(r"^NP[\s\d]", norm):
            return pat
        if pat == "INTEREST" and "INTEREST ON ACCOUNT" in norm:
            return pat
        if pat == "TAX" and norm.strip() == "TAX":
            return pat
        if pat == "ADMIN_FEE" and norm.strip() == "ADMIN FEE":
            return pat
        if pat == "MONTHLY_FEE_ATM" and norm.strip() == "MONTHLY FEE ATM":
            return pat
        if pat == "BFST" and norm.startswith("BFST"):
            return pat
        if pat == "TARIK_TUNAI" and norm.startswith("TARIK TUNAI"):
            return pat
    return UNKNOWN


# --------------------------------------------------------------------------
# Per-pattern field extraction
# --------------------------------------------------------------------------
def _extract_esb(desc: str, row: dict, notes: list[str]) -> None:
    m = ESB_RE.search(desc)
    if m:
        row["esb_system"], row["esb_code"], row["esb_reference"] = m.group(1), m.group(2), m.group(3)
    elif "ESB:" in desc.upper():
        notes.append("ESB marker present but structure not parsed")


def _parse_wbnk(desc: str, row: dict, notes: list[str]) -> None:
    m = WBNK_RE.match(desc)
    if not m:
        notes.append("WBNK prefix present but T+16-digit destination segment not matched")
        row["manual_review_required"] = True
        return
    row["internal_id"] = m.group(2)
    seg = m.group(3)
    row["raw_destination_segment"] = seg
    if seg.isdigit() and len(seg) == 16 and seg.startswith("0"):
        row["destination_account_candidate"] = seg[1:]
        row["technical_counterparty_account_candidate"] = seg[1:]
    else:
        notes.append(f"destination segment '{seg}' does not match known T0+15-digit rule")
        row["manual_review_required"] = True


def _parse_dana(desc: str, row: dict, notes: list[str]) -> None:
    row["wallet_name"] = "DANA"
    m = re.match(r"^DANA(\d{20})([A-Z ].*?)\s*/", desc)
    if m:
        row["transaction_reference"] = m.group(1)
        name = m.group(2).strip()
        if name:
            row["counterparty_name"] = name
    else:
        notes.append("DANA payload did not match <20-digit ref><name> structure")
        row["manual_review_required"] = True
    ws = re.search(r"WS_OB:(\d+);(\d+)", desc)
    if ws:
        if row["transaction_reference"] == UNKNOWN:
            row["transaction_reference"] = ws.group(1)
        row["secondary_reference"] = ws.group(2)


def _parse_prch(desc: str, row: dict, notes: list[str]) -> None:
    rest = desc[4:].strip()
    if "ESB:" in rest:
        rest = rest.split("ESB:", 1)[0].strip()
    pre, tail, secondary = rest, "", ""
    if "#" in rest:
        pre, post = rest.split("#", 1)
        m = re.match(r"(\d+)", post)
        if m:
            secondary = m.group(1)
            tail = post[m.end():]
    pre = pre.strip(" /")
    if pre.isdigit():
        row["transaction_reference"] = pre
    elif pre:
        row["merchant_name"] = pre
        tail_digits = re.findall(r"\d{6,}", tail)
        if tail_digits:
            row["transaction_reference"] = tail_digits[0]
        else:
            notes.append("merchant name found but no numeric transaction reference")
    row["secondary_reference"] = secondary or UNKNOWN


def _parse_prtt(desc: str, row: dict, notes: list[str]) -> None:
    m = re.match(r"^PRTT(\d+)#(\d+)", desc)
    if m:
        row["transaction_reference"] = m.group(1)
        row["secondary_reference"] = m.group(2)
    else:
        notes.append("PRTT payload did not match <ref>#<secondary>")
        row["manual_review_required"] = True


def _parse_np(desc: str, row: dict, notes: list[str]) -> None:
    m = re.match(r"^NP\s+([A-Z0-9]+)\s+(\d+)\s*/\s*NP(\d+)", desc)
    if m:
        row["transaction_reference"] = m.group(1)
        row["secondary_reference"] = m.group(2)
        row["long_reference_payload"] = m.group(3)
    else:
        notes.append("NP payload did not match expected structure")
        row["manual_review_required"] = True


def _parse_bfst(desc: str, row: dict, notes: list[str]) -> None:
    m = re.match(r"^BFST(\d+)(.*?)\s*/", desc)
    if m:
        row["internal_id"] = m.group(1)
        name = m.group(2).strip(" ;")
        if name:
            row["counterparty_name"] = name
    else:
        notes.append("BFST payload did not match expected structure")
        row["manual_review_required"] = True


def _parse_tarik_tunai(desc: str, row: dict, notes: list[str]) -> None:
    m = re.match(r"^TARIK TUNAI#(\d+)#(\d+)", desc)
    if m:
        row["transaction_reference"] = m.group(1)
        row["secondary_reference"] = m.group(2)
    else:
        notes.append("TARIK TUNAI payload did not match expected structure")
        row["manual_review_required"] = True


def _score_confidence(pattern: str, row: dict) -> str:
    if pattern == UNKNOWN:
        return UNKNOWN
    if row["destination_account_verified"]:
        return "HIGH"
    if pattern in {"INTEREST", "TAX", "ADMIN_FEE", "MONTHLY_FEE_ATM"}:
        return "HIGH"
    if pattern in WBNK_PATTERNS:
        return "MEDIUM" if row["destination_account_candidate"] != UNKNOWN else "LOW"
    if pattern in {"PRCH", "PRTT", "NP", "DANA", "BFST", "TARIK_TUNAI"}:
        if row["transaction_reference"] != UNKNOWN or row["merchant_name"] != UNKNOWN:
            return "MEDIUM"
        return "LOW"
    return "LOW"


# --------------------------------------------------------------------------
# Main parse
# --------------------------------------------------------------------------
def parse_transactions(
    df: pd.DataFrame,
    description_col: str = "uraian_transaksi",
    verification_db: dict[str, dict] | None = None,
    verified_at: str | None = None,
) -> pd.DataFrame:
    """Parse a DataFrame of BRI rows into a structured, evidence-annotated frame."""
    verification_db = verification_db or {}
    if description_col not in df.columns:
        raise KeyError(f"missing description column '{description_col}'")

    records: list[dict] = []
    for pos, (_, r) in enumerate(df.iterrows()):
        raw = "" if pd.isna(r.get(description_col)) else str(r.get(description_col))
        norm = normalize_description(raw)
        row = _blank_row()
        notes: list[str] = []

        pattern = detect_pattern(norm)
        row["transaction_pattern"] = pattern
        row["technical_transaction_code"] = pattern
        spec = PATTERN_REGISTRY[pattern]
        row["transaction_family"] = spec["family"]
        row["transaction_channel"] = spec["channel"]
        row["transaction_type"] = spec["type"]
        row["economic_transaction_type"] = spec["economic"]
        row["code_meaning"] = spec["code_meaning"]
        row["code_meaning_confidence"] = spec["code_meaning_confidence"]
        row["interpretation_source_type"] = spec["source_type"]
        row["evidence_level"] = spec["evidence"]
        row["account_role"] = spec["account_role"]

        debet = r.get("debet")
        kredit = r.get("kredit")
        row["transaction_direction"] = _direction(debet, kredit)

        # per-pattern extraction
        if pattern in WBNK_PATTERNS:
            _parse_wbnk(raw.upper(), row, notes)
            _extract_esb(raw, row, notes)
        elif pattern == "DANA":
            _parse_dana(raw.upper(), row, notes)
            _extract_esb(raw, row, notes)
        elif pattern == "PRCH":
            _parse_prch(raw.upper(), row, notes)
            _extract_esb(raw, row, notes)
        elif pattern == "PRTT":
            _parse_prtt(raw.upper(), row, notes)
        elif pattern == "NP":
            _parse_np(raw.upper(), row, notes)
        elif pattern == "BFST":
            _parse_bfst(raw.upper(), row, notes)
            _extract_esb(raw, row, notes)
        elif pattern == "TARIK_TUNAI":
            _parse_tarik_tunai(raw.upper(), row, notes)
            _extract_esb(raw, row, notes)
        elif pattern == UNKNOWN:
            notes.append("No recognised transaction pattern; manual review required")
            row["manual_review_required"] = True

        # account verification database
        cand = row["destination_account_candidate"]
        if cand != UNKNOWN and cand in verification_db:
            v = verification_db[cand]
            row["destination_account_verified"] = True
            row["destination_name"] = v["account_name"]
            row["destination_bank"] = v["bank_name"]
            row["evidence_level"] = VERIFIED
            notes.append(f"destination account manually verified via {v['verification_method'] or 'manual lookup'}")

        # teller hygiene
        teller = "" if pd.isna(r.get("teller")) else str(r.get("teller")).strip()
        if teller and not teller.isdigit():
            notes.append(f"teller value '{teller}' is not numeric (OCR artifact)")
            row["manual_review_required"] = True

        if row["transaction_direction"] == "REVIEW":
            notes.append("both Debet and Kredit are non-zero; direction ambiguous")
            row["manual_review_required"] = True
        if row["transaction_direction"] == "UNKNOWN":
            notes.append("neither Debet nor Kredit is populated; direction unknown")
            row["manual_review_required"] = True

        row["raw_description"] = raw
        row["normalized_description"] = norm
        row["parser_confidence"] = _score_confidence(pattern, row)
        row["parser_notes"] = "; ".join(notes) if notes else "ok"
        records.append(row)

    out = pd.DataFrame(records)

    # carry source columns through for convenience
    for col in ("tanggal_transaksi", "waktu", "teller", "debet", "kredit", "saldo"):
        if col in df.columns:
            out[col] = df[col].values

    if "tanggal_transaksi" in out.columns and "waktu" in out.columns:
        out["datetime"] = pd.to_datetime(
            out["tanggal_transaksi"].astype(str) + " " + out["waktu"].astype(str),
            format="%d/%m/%y %H:%M:%S", errors="coerce")
    out["transaction_id"] = [f"TXN{i+1:04d}" for i in range(len(out))]
    return out


# --------------------------------------------------------------------------
# Fee pairing
# --------------------------------------------------------------------------
def classify_fees(parsed: pd.DataFrame) -> pd.DataFrame:
    """Pair same-reference rows and flag the smaller debit as a service fee.

    Only WBNK (transfer family) rows are promoted to CONFIRMED/LIKELY fees.
    Other paired rows (e.g. PRTT) are only marked possible_paired_transaction
    so a human can decide.
    """
    out = parsed.copy()
    keys = [c for c in ("tanggal_transaksi", "waktu", "normalized_description", "esb_reference") if c in out.columns]
    if not keys:
        return out
    for _, idxs in out.groupby(keys, sort=False, dropna=False).groups.items():
        idxs = list(idxs)
        debits = [i for i in idxs if out.at[i, "transaction_direction"] == "OUT" and out.at[i, "debet"] > 0]
        if len(idxs) < 2 or len(debits) < 2:
            continue
        main_i = max(debits, key=lambda i: out.at[i, "debet"])
        for i in idxs:
            out.at[i, "possible_paired_transaction"] = True
        main_amount = out.at[main_i, "debet"]
        for i in debits:
            if i == main_i:
                continue
            ratio = out.at[i, "debet"] / main_amount if main_amount else 1.0
            if ratio >= 0.5:
                continue
            out.at[i, "fee_amount"] = int(out.at[i, "debet"])
            out.at[i, "parent_transaction_id"] = out.at[main_i, "transaction_id"]
            if out.at[i, "transaction_pattern"] in WBNK_PATTERNS:
                same_teller = str(out.at[i, "teller"]) == str(out.at[main_i, "teller"])
                out.at[i, "is_fee_transaction"] = True
                out.at[i, "fee_classification"] = "CONFIRMED" if same_teller else "LIKELY"
                out.at[i, "parser_notes"] = (str(out.at[i, "parser_notes"]) +
                                             f"; service fee paired with {out.at[main_i,'transaction_id']}"
                                             f" (ratio {ratio:.4f}, teller {'match' if same_teller else 'differs'})")
            else:
                out.at[i, "fee_classification"] = UNKNOWN
                out.at[i, "manual_review_required"] = True
                out.at[i, "parser_notes"] = (str(out.at[i, "parser_notes"]) +
                                             "; smaller debit shares reference with a larger row but is NOT confirmed as a fee")
    return out


# --------------------------------------------------------------------------
# Interest / tax relationship
# --------------------------------------------------------------------------
def link_interest_tax(parsed: pd.DataFrame) -> pd.DataFrame:
    out = parsed.copy()
    for i in out.index:
        if out.at[i, "transaction_pattern"] != "TAX":
            continue
        if i == 0 or out.at[i - 1, "transaction_pattern"] != "INTEREST":
            continue
        interest = out.at[i - 1, "kredit"]
        tax = out.at[i, "debet"]
        if pd.isna(interest) or pd.isna(tax) or interest in (0, None):
            continue
        rate = float(tax) / float(interest)
        out.at[i, "tax_relationship_candidate"] = "INTEREST_TAX"
        out.at[i, "related_interest_amount"] = int(interest)
        out.at[i, "effective_tax_rate"] = round(rate, 6)
        out.at[i, "tax_relationship_confidence"] = "HIGH" if abs(rate - 0.20) < 0.005 else "MEDIUM"
        out.at[i, "parser_notes"] = (str(out.at[i, "parser_notes"]) +
                                     f"; follows interest row {out.at[i-1,'transaction_id']} (rate {rate:.2%})")
    return out


# --------------------------------------------------------------------------
# Reporting sheets
# --------------------------------------------------------------------------
BANK_INTERNAL_PATTERNS = {"INTEREST", "TAX", "ADMIN_FEE", "MONTHLY_FEE_ATM"}

def _running_balance(parsed: pd.DataFrame) -> pd.DataFrame:
    """Attach per-row opening balance (balance before each txn), expected balance, difference and match flag.

    For row i, expected balance = (balance before row i) - debit_i + credit_i.
    Balance before row 0 is the global opening balance = saldo_0 + debit_0 - credit_0.
    Balance before row i>0 is the actual saldo of row i-1.
    This matches the validation in src/manual_loader.validate_running_balance.
    """
    out = parsed.copy()
    debet = out["debet"].fillna(0).astype("int64") if "debet" in out.columns else pd.Series(0, index=out.index, dtype="int64")
    kredit = out["kredit"].fillna(0).astype("int64") if "kredit" in out.columns else pd.Series(0, index=out.index, dtype="int64")
    saldo = out["saldo"].astype("int64") if "saldo" in out.columns else pd.Series([], dtype="int64")

    if len(out) == 0:
        out["opening_balance"] = 0
        out["expected_saldo"] = []
        out["saldo_difference"] = []
        out["reconciliation_status"] = []
        return out

    # global opening = saldo of first row + debet_0 - kredit_0
    global_open = int(saldo.iloc[0]) + int(debet.iloc[0]) - int(kredit.iloc[0])

    prev_bal_list = []
    expected = []
    prev = global_open
    for i in out.index:
        d = int(debet.loc[i])
        k = int(kredit.loc[i])
        sv = int(saldo.loc[i])
        exp = prev - d + k
        expected.append(exp)
        prev_bal_list.append(prev)
        prev = sv
    out["opening_balance"] = prev_bal_list  # per-row opening (balance before this txn)
    out["expected_saldo"] = expected
    out["saldo_difference"] = out["saldo"].astype("int64") - pd.Series(expected, index=out.index)
    out["reconciliation_status"] = out["saldo_difference"].apply(lambda x: "MATCH" if x == 0 else "MISMATCH")
    return out


def _payee_identity(row: pd.Series) -> tuple[str, str]:
    """Return (destination_type, payee_label) without fabricating identities."""
    acct = str(row.get("destination_account_candidate", UNKNOWN))
    if acct and acct != UNKNOWN and not pd.isna(acct):
        name = row.get("destination_name", UNKNOWN)
        label = f"{name} ({acct})" if name and name != UNKNOWN else acct
        return "Account Transfer", label
    merchant = str(row.get("merchant_name", UNKNOWN))
    if merchant and merchant != UNKNOWN:
        return "Merchant", merchant
    wallet = str(row.get("wallet_name", UNKNOWN))
    if wallet and wallet != UNKNOWN:
        return "E-Wallet", wallet
    cp = str(row.get("counterparty_name", UNKNOWN))
    if cp and cp != UNKNOWN:
        return "Counterparty", cp
    pattern = row.get("transaction_pattern", UNKNOWN)
    if pattern == "TARIK_TUNAI":
        return "Cash Withdrawal", "ATM (TARIK TUNAI)"
    if pattern in {"PRCH", "PRTT", "NP"}:
        ref = row.get("transaction_reference", UNKNOWN)
        suffix = f" [{ref}]" if ref and ref != UNKNOWN else ""
        return "Merchant", f"Merchant not stated in description ({pattern}){suffix}"
    if pattern in BANK_INTERNAL_PATTERNS:
        return "Bank System", str(pattern)
    return "Other", str(pattern)


def build_sheets(parsed: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Build the single-workbook tabs: list, summary, balance control, destination."""
    sheets: dict[str, pd.DataFrame] = {}
    parsed = _running_balance(parsed)

    # ---- Tab 1: full transaction list (source order) -----------------------
    tx = parsed.copy()
    tx["_payee"] = tx.apply(_payee_identity, axis=1)
    tx["destination_type"] = [t for t, _ in tx["_payee"]]
    tx["destination_label"] = [l for _, l in tx["_payee"]]

    def _amount(v):
        try:
            return int(v) if not pd.isna(v) else 0
        except (TypeError, ValueError):
            return 0

    list_rows = []
    for i, r in tx.iterrows():
        list_rows.append({
            "No": len(list_rows) + 1,
            "Transaction ID": r.get("transaction_id", ""),
            "Date": r.get("datetime"),
            "Tanggal": r.get("tanggal_transaksi", ""),
            "Waktu": r.get("waktu", ""),
            "Uraian Transaksi": r.get("raw_description", ""),
            "Jenis Transaksi": r.get("transaction_type", UNKNOWN),
            "Kategori": r.get("transaction_pattern", UNKNOWN),
            "Arah": "DEBIT (keluar)" if r.get("transaction_direction") == "OUT"
                    else "KREDIT (masuk)" if r.get("transaction_direction") == "IN" else UNKNOWN,
            "Debit": _amount(r.get("debet")),
            "Credit": _amount(r.get("kredit")),
            "Saldo": _amount(r.get("saldo")),
            "Saldo Seharusnya": _amount(r.get("expected_saldo")),
            "Selisih": _amount(r.get("saldo_difference")),
            "Status Rekonsiliasi": r.get("reconciliation_status", UNKNOWN),
            "Teller": r.get("teller", ""),
            "Tujuan/ Merchant": r.get("destination_label", UNKNOWN),
            "Tipe Tujuan": r.get("destination_type", UNKNOWN),
            "Rekening Tujuan": r.get("destination_account_candidate", UNKNOWN),
            "Nama Tujuan": r.get("destination_name", UNKNOWN),
            "Bank Tujuan": r.get("destination_bank", UNKNOWN),
            "Referensi": r.get("transaction_reference", UNKNOWN),
            "ESB Reference": r.get("esb_reference", UNKNOWN),
            "Biaya (Fee)": _amount(r.get("fee_amount")),
            "Klasifikasi Fee": r.get("fee_classification", UNKNOWN),
            "Transaksi Induk": r.get("parent_transaction_id", ""),
            "Tingkat Keyakinan": r.get("parser_confidence", UNKNOWN),
            "Perlu Review Manual": bool(r.get("manual_review_required", False)),
            "Catatan Parser": r.get("parser_notes", ""),
        })
    tx_sheet = pd.DataFrame(list_rows)
    sheets["Daftar Transaksi"] = tx_sheet

    # ---- Tab 2: Ringkasan ---------------------------------------------------
    total_debit = int(tx["debet"].fillna(0).sum())
    total_kredit = int(tx["kredit"].fillna(0).sum())
    opening = int(tx["opening_balance"].iloc[0]) if len(tx) else 0
    closing = int(tx["saldo"].iloc[-1]) if len(tx) else 0
    mismatches = int((tx["reconciliation_status"] != "MATCH").sum())
    dmin = tx["datetime"].min()
    dmax = tx["datetime"].max()
    summary_rows = [
        ("Sumber data", "manual.txt (satu-satunya sumber)"),
        ("Periode", f"{dmin} — {dmax}"),
        ("Jumlah transaksi", len(tx)),
        ("Total Debit", total_debit),
        ("Total Kredit", total_kredit),
        ("Saldo Awal", opening),
        ("Saldo Akhir", closing),
        ("Selisih (Awal - Debit + Kredit - Akhir)", opening - total_debit + total_kredit - closing),
        ("Jumlah baris selisih != 0", mismatches),
        ("Status rekonsiliasi", "COCOK (0 selisih)" if mismatches == 0 else f"TIDAK COCOK ({mismatches} baris)"),
        ("Jumlah transaksi fee teridentifikasi", int(tx["is_fee_transaction"].sum())),
        ("Jumlah transaksi perlu review manual", int(tx["manual_review_required"].sum())),
    ]
    sheets["Ringkasan"] = pd.DataFrame(summary_rows, columns=["Keterangan", "Nilai"])

    # ---- Tab 3: Kontrol Saldo (per baris + batas bulan) ---------------------
    ctrl = tx[["transaction_id", "datetime", "raw_description", "debet", "kredit",
               "opening_balance", "expected_saldo", "saldo", "saldo_difference",
               "reconciliation_status"]].copy()
    ctrl = ctrl.rename(columns={
        "transaction_id": "Transaction ID", "datetime": "Date", "raw_description": "Uraian Transaksi",
        "debet": "Debit", "kredit": "Credit", "opening_balance": "Saldo Awal",
        "expected_saldo": "Saldo Seharusnya", "saldo": "Saldo Aktual",
        "saldo_difference": "Selisih", "reconciliation_status": "Status",
    })
    sheets["Kontrol Saldo"] = ctrl

    # ---- Tab 4: Rekap Bulanan ----------------------------------------------
    tx["_month"] = pd.to_datetime(tx["datetime"]).dt.strftime("%Y-%m")
    monthly_rows = []
    for month, g in tx.groupby("_month", sort=True):
        month_open = int(g["opening_balance"].iloc[0])
        month_debit = int(g["debet"].fillna(0).sum())
        month_credit = int(g["kredit"].fillna(0).sum())
        month_close = int(g["saldo"].iloc[-1])
        monthly_rows.append({
            "Bulan": month,
            "Jumlah Transaksi": len(g),
            "Saldo Awal": month_open,
            "Total Debit": month_debit,
            "Total Kredit": month_credit,
            "Saldo Akhir": month_close,
            "Selisih": month_open - month_debit + month_credit - month_close,
            "Status": "COCOK" if month_open - month_debit + month_credit - month_close == 0 else "TIDAK COCOK",
        })
    sheets["Rekap Bulanan"] = pd.DataFrame(monthly_rows)

    # ---- Tab 5: Destination Summary (transfer + merchant + e-wallet) -------
    dest = tx[(~tx["is_fee_transaction"]) & (~tx["transaction_pattern"].isin(BANK_INTERNAL_PATTERNS))].copy()
    dest = dest[dest["transaction_direction"] == "OUT"]
    rows = []
    for (dtype, label), g in dest.groupby(["destination_type", "destination_label"], sort=False):
        fees = tx[(tx["parent_transaction_id"].isin(g["transaction_id"])) & (tx["is_fee_transaction"])]
        verified = bool(g["destination_account_verified"].any()) if "destination_account_verified" in g else False
        rows.append({
            "Tipe Tujuan": dtype,
            "Tujuan / Merchant": label,
            "Rekening Tujuan": g["destination_account_candidate"].iloc[0],
            "Bank": g["destination_bank"].iloc[0],
            "Status Verifikasi": "VERIFIED" if verified else "UNVERIFIED",
            "Jumlah Transaksi": int(len(g)),
            "Total Nominal": int(g["debet"].fillna(0).sum()),
            "Rata-rata": int(round(g["debet"].fillna(0).mean())) if len(g) else 0,
            "Transaksi Terbesar": int(g["debet"].fillna(0).max()) if len(g) else 0,
            "Transaksi Pertama": str(g["tanggal_transaksi"].iloc[0]),
            "Transaksi Terakhir": str(g["tanggal_transaksi"].iloc[-1]),
            "Total Fee Terkait": int(fees["fee_amount"].fillna(0).sum()) if len(fees) else 0,
            "Referensi": "; ".join(str(x) for x in g["esb_reference"].tolist() if str(x) not in ("", UNKNOWN)),
        })
    dest_cols = ["Tipe Tujuan", "Tujuan / Merchant", "Rekening Tujuan", "Bank", "Status Verifikasi",
                 "Jumlah Transaksi", "Total Nominal", "Rata-rata", "Transaksi Terbesar",
                 "Transaksi Pertama", "Transaksi Terakhir", "Total Fee Terkait", "Referensi"]
    sheets["Destination Summary"] = pd.DataFrame(rows, columns=dest_cols) if rows else pd.DataFrame(columns=dest_cols)

    # ---- Tab 6: Ringkasan Jenis Transaksi ----------------------------------
    total = len(tx)
    rows = []
    for pat in PATTERN_PRIORITY:
        g = tx[tx["transaction_pattern"] == pat]
        if g.empty:
            continue
        rows.append({
            "Jenis Transaksi": pat,
            "Tipe": PATTERN_REGISTRY[pat]["type"],
            "Kanal": PATTERN_REGISTRY[pat]["channel"],
            "Jumlah": int(len(g)),
            "Total Debit": int(g["debet"].fillna(0).sum()),
            "Total Kredit": int(g["kredit"].fillna(0).sum()),
            "Persentase Transaksi": len(g) / total if total else 0,
        })
    sheets["Ringkasan Jenis Transaksi"] = pd.DataFrame(rows)

    return sheets



# --------------------------------------------------------------------------
# Workbook formatting by semantic column type
# --------------------------------------------------------------------------
CURRENCY_FORMAT = '"Rp" #,##0'
PERCENT_FORMAT = "0.00%"
INTEGER_FORMAT = "#,##0"
TEXT_FORMAT = "@"
DATETIME_FORMAT = "dd/mm/yyyy hh:mm:ss"

CURRENCY_HEADERS = {
    "Debit", "Credit", "Saldo", "Saldo Seharusnya", "Saldo Aktual", "Saldo Awal", "Saldo Akhir",
    "Fee Amount", "Biaya (Fee)", "Total Nominal", "Rata-rata", "Transaksi Terbesar",
    "Total Fee Terkait", "Total Transfer", "Average Transfer", "Highest Transfer",
    "Total Related Fee", "Total Debit", "Total Credit", "Total Main Debit", "Total Fee",
    "Related Interest", "Selisih", "Amount",
}
PERCENT_HEADERS = {"Effective Tax Rate", "Percent of Transactions", "Persentase Transaksi"}
DATE_HEADERS = {"Date"}
INTEGER_HEADERS = {
    "No", "Jumlah Transaksi", "Jumlah", "Transaction Count", "WBNKTRIK Count",
    "WBNKTRF Count", "Unique ESB References", "Count",
}
BOOLEAN_HEADERS = {"Destination Verified", "Is Fee", "Possible Paired", "Manual Review", "Perlu Review Manual"}
TEXT_HEADERS = {"Teller", "Transaction ID", "Tanggal", "Waktu"}


def _column_kind(header: str) -> str:
    if header in CURRENCY_HEADERS:
        return "currency"
    if header in PERCENT_HEADERS:
        return "percent"
    if header in DATE_HEADERS:
        return "date"
    if header in INTEGER_HEADERS:
        return "integer"
    if header in BOOLEAN_HEADERS:
        return "boolean"
    return "text"


def write_workbook(sheets: dict[str, pd.DataFrame], path: Path) -> None:
    """Write workbook with currency, percentage, count, date, boolean and text formats."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        try:
            path.unlink()
        except PermissionError as exc:
            raise PermissionError(f"{path} is locked (open in Excel?): {exc}") from exc
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, data in sheets.items():
            data.to_excel(writer, sheet_name=name[:31], index=False)
    wb = load_workbook(path)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [cell.value for cell in ws[1]]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, col_cells in enumerate(ws.iter_cols(min_row=2, max_row=ws.max_row), start=1):
            header = str(headers[col_idx - 1] or "")
            kind = _column_kind(header)
            for cell in col_cells:
                if kind == "currency" and isinstance(cell.value, (int, float)):
                    cell.number_format = CURRENCY_FORMAT
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif kind == "percent" and isinstance(cell.value, (int, float)):
                    cell.number_format = PERCENT_FORMAT
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif kind == "integer" and isinstance(cell.value, (int, float)):
                    cell.number_format = INTEGER_FORMAT
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif kind == "date" and cell.value is not None:
                    cell.number_format = DATETIME_FORMAT
                    cell.alignment = Alignment(horizontal="left", vertical="top")
                elif kind == "boolean":
                    cell.number_format = "General"
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    cell.number_format = TEXT_FORMAT
                    cell.alignment = Alignment(
                        horizontal="left", vertical="top",
                        wrap_text=(header in {"Uraian Transaksi", "Catatan Parser", "Referensi"}),
                    )

        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = max((len(str(c.value or "")) for c in col), default=0) + 3
            ws.column_dimensions[letter].width = min(65, max(12, width))

    wb.save(path)
